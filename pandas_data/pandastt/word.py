from docx import Document
import pandas  as pd
from xlwt import *
import re


def get_text(text):
    text = text.strip().replace(" ", "").replace(u'\u3000', u' ').replace(u'\xa0', u' ')
    return text


def is_number(s):
    try:
        float(s)
        return True
    except ValueError:
        return False


class Word(object):
    def __init__(self, path=''):
        # 获得文档
        self.file = Document("resource/question1.docx")

    def start(self):
        que_num = 0
        que_item = None
        que_list = []
        # 按照段落读取文档内容
        for para in self.file.paragraphs:
            if not para.text: continue
            # if que_num==135:
            #     text=para.text
            #     print(que_num,text)
            if self.is_que_title(para.text):
                que_num = que_num + 1
                if que_item:
                    self.config_que(que_item)
                    que_list.append(que_item)
                que_item = {}
                que_item['num'] = que_num
                que_item['text'] = get_text(para.text)
                que_item['ans'] = self.get_que_ans(para, que_item)
            else:
                que_item['text'] = que_item['text'] + get_text(para.text)
                que_item['ans'] = self.get_que_ans(para, que_item)
            # print(que_item)
        # for que in que_list:
        # print(que)
        # if que and "ans" in que and len(que['ans']) <= 0:
        #     print(que)
        # if que and "options" in que and len(que['options']) <= 0:
        #     print(que)
        # if que and "title" in que and len(que['title']) <= 0:
        #     print(que)
        self.save_to_excel(que_list)

    def save_to_excel(self, que_list):
        que_obj = ('num', 'title', 'ans', 'options')
        file = Workbook(encoding='utf-8')
        table = file.add_sheet('data')
        table.col(0).width = 1500
        table.col(1).width = 15000
        table.col(2).width = 9000
        table.col(3).width = 15000
        table.col(4).width = 15000
        for i, que in enumerate(que_list):
            for j, key in enumerate(que_obj):
                if i == 0:
                    value = key
                elif key in que:
                    value = que[key]
                else:
                    value = ''
                alignment = Alignment()
                if i == 0 or j == 0:
                    alignment.horz = Alignment.HORZ_CENTER
                else:
                    alignment.horz = Alignment.HORZ_LEFT
                style = XFStyle()
                style.alignment = alignment
                table.write(i, j, value, style)
        file.save('data.xlsx')

    def is_que_title(self, text):
        texts = text.split("、")
        if texts and text[0] == "一":
            return True
        if len(texts) > 0 and is_number(text[0]):
            return True
        else:
            return False

    def get_que_ans(self, para, que_item):
        flag = -1
        # 1.判断被下划线标识的字段
        runs = para.runs
        if "ans" in que_item:
            ans = que_item['ans']
        else:
            ans = ""
        for run in runs:
            if run.font.underline:
                ans = ans + get_text(run.text)
                flag = 1
        # 2.判断被括号包围的字段
        if "（" in para.text and "）" in para.text and flag == -1:
            an = para.text.split("（")[1].split("）")[0]
            if "A" in an or "B" in an or "C" in an or "D" in an:
                ans = ans + an
                flag = 2
        if "(" in para.text and ")" in para.text and flag == -1:
            an = para.text.split("(")[1].split(")")[0]
            if "A" in an or "B" in an or "C" in an or "D" in an:
                ans = ans + an
                flag = 3
        return ans

    def config_que(self, que_item):
        if "。" in que_item['text']:
            que_item['title'] = que_item['text'].split("。")[0]
            que_item['options'] = que_item['text'].split("。")[1]
        elif "）：" in que_item['text']:
            que_item['title'] = que_item['text'].split("）：")[0] + "）："
            que_item['options'] = que_item['text'].split("）：")[1]
        elif "A." in que_item['text']:
            que_item['title'] = que_item['text'].split("A.")[0]
            que_item['options'] = que_item['text'].split("A.")[1] + "A."
        elif "A．" in que_item['text']:
            que_item['title'] = que_item['text'].split("A．")[0]
            que_item['options'] = que_item['text'].split("A．")[1] + "A."
        elif "A、" in que_item['text']:
            que_item['title'] = que_item['text'].split("A、")[0]
            que_item['options'] = que_item['text'].split("A、")[1] + "A."
        if 'title' in que_item:
            que_item['title'] = re.sub(r'\（(.+?)\）', "()", que_item['title'])
            que_item['title'] = re.sub(r'\((.+?)\)', "()", que_item['title'])
        if 'ans' in que_item:
            print(que_item['ans'])
            ans = []
            if "A" in que_item['ans']:
                ans.append("A")
            if "B" in que_item['ans']:
                ans.append("B")
            if "C" in que_item['ans']:
                ans.append("C")
            if "D" in que_item['ans']:
                ans.append("D")
            que_item['ans'] = ans


if __name__ == '__main__':
    word = Word()
    word.start()
